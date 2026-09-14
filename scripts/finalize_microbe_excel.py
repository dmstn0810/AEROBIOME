import sys
from pathlib import Path
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import shutil

sys.stdout.reconfigure(encoding='utf-8')
sys.path.append(str(Path.cwd()))

orig_file = Path("괴산연습림 전체데이터.xlsx")
backup_file = Path("괴산연습림 전체데이터_원본백업.xlsx")
if not backup_file.exists():
    shutil.copy(orig_file, backup_file)
    print(f"Created backup: {backup_file}")

from scripts.process_microbe_dataset import df

# Let's inspect df columns
# df currently has:
# ['site no.', 'date', 'time', 'vege', 'air-temp', 'air-RH', 'windspeed', 'PM10', 'PM2.5', 'illum',
#  'soil-PH', 'soil-temp', 'soil-RH', 'PCA-B', 'PDA-F', 'n-ion', 'rep_idx', 'PCA-B_CFU', 'PDA-F_CFU']

# The user requested:
# 1. "공기미생물 농도값은 PCA-B와 PDA-F이며 현재는 콜로니개수이기 때문에 CFU/m^3으로 바꾸면 콜로니개수/0.1을 해야해"
# 2. "또한 중간에 n이라는 null값이 있으니 이것도 채워줘"
# 3. ""괴산연습림 전체데이터"는 음이온 값이 비어있어서 각 조사날짜별 조사대상지에 대표값 5개가 필요해"

# Create updated workbook
wb = openpyxl.Workbook()
wb.remove(wb.active) # remove default

# Styles
header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
accent_fill = PatternFill(start_color="276A3C", end_color="276A3C", fill_type="solid")
header_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
title_font = Font(name="맑은 고딕", size=15, bold=True, color="1F497D")
sub_font = Font(name="맑은 고딕", size=11, bold=True, color="2C5E8A")
bold_font = Font(name="맑은 고딕", size=9, bold=True)
regular_font = Font(name="맑은 고딕", size=9)
thin_side = Side(border_style="thin", color="D9D9D9")
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
center_align = Alignment(horizontal="center", vertical="center")
right_align = Alignment(horizontal="right", vertical="center")
left_align = Alignment(horizontal="left", vertical="center")

# -------------------------------------------------------------
# Sheet 1: 원본 (Updated as requested)
# -------------------------------------------------------------
ws1 = wb.create_sheet(title="원본")
ws1.views.sheetView[0].showGridLines = True

# Headers matching original + clearly indicating CFU/m^3
orig_headers = [
    'site no.', 'date', 'time', 'vege', 'air-temp', 'air-RH', 'windspeed', 'PM10', 'PM2.5', 'illum',
    'soil-PH', 'soil-temp', 'soil-RH', 'PCA-B (CFU/m3)', 'PDA-F (CFU/m3)', 'n-ion (개/cm3)',
    'PCA-B (콜로니수)', 'PDA-F (콜로니수)'
]

for c_i, h in enumerate(orig_headers, 1):
    c = ws1.cell(1, c_i, h)
    c.fill = header_fill if c_i <= 16 else PatternFill(start_color="595959", fill_type="solid")
    c.font = header_font
    c.alignment = center_align
    c.border = thin_border

for r_i, (_, r) in enumerate(df.iterrows(), 2):
    d_val = r['date']
    t_val = r['time']
    
    # row values
    row_data = [
        int(r['site no.']),
        str(d_val)[:10],
        str(t_val),
        r['vege'],
        round(float(r['air-temp']), 2),
        round(float(r['air-RH']), 2),
        round(float(r['windspeed']), 2),
        round(float(r['PM10']), 2),
        round(float(r['PM2.5']), 2),
        round(float(r['illum']), 1),
        round(float(r['soil-PH']), 2),
        round(float(r['soil-temp']), 1),
        round(float(r['soil-RH']), 1),
        round(float(r['PCA-B_CFU']), 1), # CFU/m3
        round(float(r['PDA-F_CFU']), 1), # CFU/m3
        round(float(r['n-ion']), 1) if pd.notna(r['n-ion']) else "", # n-ion
        round(float(r['PCA-B']), 1),     # original colony count
        round(float(r['PDA-F']), 1)      # original colony count
    ]
    
    for c_i, val in enumerate(row_data, 1):
        cell = ws1.cell(r_i, c_i, val)
        cell.font = regular_font
        cell.border = thin_border
        if c_i in [1, 2, 3]:
            cell.alignment = center_align
        elif c_i == 4:
            cell.alignment = left_align
        else:
            cell.alignment = right_align

ws1.auto_filter.ref = f"A1:R{len(df) + 1}"

# Auto-adjust column widths
for col in ws1.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws1.column_dimensions[col_letter].width = max(max_len + 3, 11)

# -------------------------------------------------------------
# Sheet 2: 정제및단위설명 (Data Dictionary & Processing Notes)
# -------------------------------------------------------------
ws2 = wb.create_sheet(title="정제및단위설명")
ws2.views.sheetView[0].showGridLines = True

ws2.cell(2, 2, "괴산연습림 데이터 정제, 단위 변환 및 음이온 대표값 연계 설명서").font = title_font

notes_data = [
    ("항목", "처리 내용 및 산출 기준"),
    ("대조구 기준", "사용자 지침에 따라 기준 대조구는 '밭(대조구)'(Site 4)을 적용함. 주차장(대조구)(Site 5)은 참고용으로 분류."),
    ("음이온(n-ion) 대표값 5개", "각 조사일자별 10개 지점의 900초(15분) 1초 연속측정 데이터를 5개 연속 3분 구간(1~3분, 4~6분, 7~9분, 10~12분, 13~15분, 각 180초)으로 등분하여 각 구간 평균값 5개를 5개 반복 행에 순차 매핑함. (5개 대표값의 평균은 해당 15분 세션의 전체 평균과 정확히 일치)"),
    ("공기미생물 농도(CFU/m3) 변환", "PCA-B(세균) 및 PDA-F(진균)의 원본 콜로니 개수를 에어샘플러 흡인량(0.1 m3 = 100 L) 기준에 맞추어 [콜로니 개수 / 0.1] = [콜로니 개수 × 10] CFU/m3로 변환 완료. 원본 콜로니 개수는 Q, R열에 보존."),
    ("결측치('n') 대체(Imputation)", "1) 기상 및 환경요인(air-temp, air-RH, windspeed, PM10, PM2.5, illum): 1~3회차 실측값의 지점별 평균으로 4~5회차 'n' 대체.\n2) 미생물(PCA-B, PDA-F): 실측된 반복구 평균으로 'n' 대체.\n3) 주차장(Site 5) 토양요인(soil-PH, soil-temp, soil-RH): 주차장은 아스팔트 포장면으로 토양 측정이 불가능하므로, 기준 대조구인 밭(Site 4)의 당일 토양 평균값으로 연계 보완."),
    ("음이온 미측정 일자 안내", "음이온 현장 조사는 2026-04-15, 04-29, 07-27, 08-13 총 4개 일자에 수행되었으며, 2026-03-04와 05-13은 음이온 측정이 미실시되어 공란으로 보존함.")
]

r_cur = 4
for item, desc in notes_data:
    c1 = ws2.cell(r_cur, 2, item)
    c2 = ws2.cell(r_cur, 3, desc)
    c1.border = thin_border
    c2.border = thin_border
    if r_cur == 4:
        c1.fill = header_fill; c1.font = header_font; c1.alignment = center_align
        c2.fill = header_fill; c2.font = header_font; c2.alignment = center_align
    else:
        c1.fill = PatternFill(start_color="F2F2F2", fill_type="solid")
        c1.font = bold_font; c1.alignment = left_align
        c2.font = regular_font; c2.alignment = left_align
    r_cur += 1

ws2.column_dimensions['B'].width = 28
ws2.column_dimensions['C'].width = 90

# -------------------------------------------------------------
# Sheet 3: 대조구(밭)_산림_비교분석 (Forest vs Field Control Table)
# -------------------------------------------------------------
ws3 = wb.create_sheet(title="대조구(밭)_산림_연계비교")
ws3.views.sheetView[0].showGridLines = True

ws3.cell(2, 2, "산림(8개 지점) vs 밭(대조구) 공기미생물 및 음이온 시기별 비교표").font = title_font

tbl_headers = [
    "조사일자", "구분", "세균 농도 (PCA-B, CFU/m3)", "세균 비율 (산림/밭)",
    "진균 농도 (PDA-F, CFU/m3)", "진균 비율 (산림/밭)",
    "음이온 농도 (n-ion, 개/cm3)", "음이온 배율 (산림/밭)"
]
r_cur = 4
for c_i, h in enumerate(tbl_headers, 2):
    c = ws3.cell(r_cur, c_i, h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center_align
    c.border = thin_border

from scripts.analyze_microbe_ion_linkage import fc_df

r_cur = 5
for _, row in fc_df.iterrows():
    # Forest row
    d_str = row['date']
    
    # Forest
    ws3.cell(r_cur, 2, d_str).alignment = center_align
    ws3.cell(r_cur, 3, "산림(8지점 평균)").alignment = left_align
    ws3.cell(r_cur, 4, round(row['forest_bact_mean'], 1)).alignment = right_align
    ws3.cell(r_cur, 5, f"{row['bact_ratio(F/C)']:.2f}배 ({row['bact_ratio(F/C)']-1:+.0%})").alignment = center_align
    ws3.cell(r_cur, 6, round(row['forest_fungi_mean'], 1)).alignment = right_align
    ws3.cell(r_cur, 7, f"{row['fungi_ratio(F/C)']:.2f}배 ({row['fungi_ratio(F/C)']-1:+.0%})").alignment = center_align
    ws3.cell(r_cur, 8, round(row['forest_ion_mean'], 1)).alignment = right_align
    ws3.cell(r_cur, 9, f"{row['ion_ratio(F/C)']:.2f}배 ({row['ion_ratio(F/C)']-1:+.0%})").alignment = center_align
    
    for ci in range(2, 10):
        ws3.cell(r_cur, ci).font = regular_font
        ws3.cell(r_cur, ci).border = thin_border
    r_cur += 1
    
    # Field control row
    ws3.cell(r_cur, 2, d_str).alignment = center_align
    ws3.cell(r_cur, 3, "대조구(밭)").alignment = left_align
    ws3.cell(r_cur, 4, round(row['field_bact_mean'], 1)).alignment = right_align
    ws3.cell(r_cur, 5, "1.00배 (기준)").alignment = center_align
    ws3.cell(r_cur, 6, round(row['field_fungi_mean'], 1)).alignment = right_align
    ws3.cell(r_cur, 7, "1.00배 (기준)").alignment = center_align
    ws3.cell(r_cur, 8, round(row['field_ion_mean'], 1)).alignment = right_align
    ws3.cell(r_cur, 9, "1.00배 (기준)").alignment = center_align
    
    for ci in range(2, 10):
        ws3.cell(r_cur, ci).font = regular_font
        ws3.cell(r_cur, ci).border = thin_border
        ws3.cell(r_cur, ci).fill = PatternFill(start_color="FFF2CC", fill_type="solid")
    r_cur += 1

# Auto adjust widths
for col_letter in ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
    ws3.column_dimensions[col_letter].width = 22

# -------------------------------------------------------------
# Sheet 4: 수종별_상세프로파일 (Species Profiles)
# -------------------------------------------------------------
ws4 = wb.create_sheet(title="수종별_상세프로파일")
ws4.views.sheetView[0].showGridLines = True

ws4.cell(2, 2, "수종(임상)별 공기미생물 및 음이온 농도 프로파일").font = title_font

from scripts.analyze_microbe_ion_linkage import sp_summary

sp_headers = ["수종분류", "표본수(N)", "세균 평균 (CFU/m3)", "세균 표준편차", "진균 평균 (CFU/m3)", "진균 표준편차", "음이온 평균 (개/cm3)", "음이온 표준편차", "평균기온(℃)", "평균습도(%)"]
r_cur = 4
for c_i, h in enumerate(sp_headers, 2):
    c = ws4.cell(r_cur, c_i, h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = center_align
    c.border = thin_border

r_cur = 5
for sp_name, row in sp_summary.iterrows():
    vals = [sp_name, int(row['n']), row['bact_mean'], row['bact_sd'], row['fungi_mean'], row['fungi_sd'], row['ion_mean'], row['ion_sd'], row['temp_mean'], row['rh_mean']]
    for c_i, val in enumerate(vals, 2):
        c = ws4.cell(r_cur, c_i, val)
        c.font = regular_font
        c.border = thin_border
        if c_i == 2:
            c.alignment = left_align
            c.font = bold_font
        elif c_i == 3:
            c.alignment = center_align
        else:
            c.alignment = right_align
            c.number_format = "#,##0.0"
    r_cur += 1

for col_letter in ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K']:
    ws4.column_dimensions[col_letter].width = 18

# -------------------------------------------------------------
# Sheet 5: 환경및미생물_상관행렬 (Correlation Matrix)
# -------------------------------------------------------------
ws5 = wb.create_sheet(title="환경_음이온_상관분석")
ws5.views.sheetView[0].showGridLines = True

ws5.cell(2, 2, "음이온, 공기미생물 및 환경요인 간 피어슨 상관계수(r) 행렬").font = title_font

from scripts.analyze_microbe_ion_linkage import corr_matrix, env_cols

kor_map = {
    'PCA-B_CFU': '세균(CFU/m3)', 'PDA-F_CFU': '진균(CFU/m3)', 'n-ion': '음이온(개/cm3)',
    'air-temp': '대기온도', 'air-RH': '대기습도', 'windspeed': '풍속',
    'PM10': '미세먼지(PM10)', 'PM2.5': '초미세먼지(PM2.5)', 'illum': '조도',
    'soil-PH': '토양pH', 'soil-temp': '토양온도', 'soil-RH': '토양습도'
}

r_cur = 4
ws5.cell(r_cur, 2, "변수명").fill = header_fill; ws5.cell(r_cur, 2).font = header_font; ws5.cell(r_cur, 2).border = thin_border
for c_i, col in enumerate(env_cols, 3):
    c = ws5.cell(r_cur, c_i, kor_map[col])
    c.fill = header_fill; c.font = header_font; c.alignment = center_align; c.border = thin_border

r_cur = 5
for row_var in env_cols:
    ws5.cell(r_cur, 2, kor_map[row_var]).font = bold_font; ws5.cell(r_cur, 2).border = thin_border; ws5.cell(r_cur, 2).alignment = left_align
    for c_i, col_var in enumerate(env_cols, 3):
        val = corr_matrix.loc[row_var, col_var]
        c = ws5.cell(r_cur, c_i, round(val, 3))
        c.font = regular_font; c.border = thin_border; c.alignment = right_align; c.number_format = "+0.000;-0.000;0.000"
        if abs(val) > 0.4 and row_var != col_var:
            c.fill = PatternFill(start_color="D9E1F2", fill_type="solid")
            c.font = bold_font
    r_cur += 1

ws5.column_dimensions['B'].width = 18
for c_i in range(3, 3 + len(env_cols)):
    col_letter = get_column_letter(c_i)
    ws5.column_dimensions[col_letter].width = 16

# Save files:
# 1. Update 괴산연습림 전체데이터.xlsx
wb.save(orig_file)
print(f"Updated {orig_file} successfully!")

# 2. Also save to outputs/
out_report = Path("outputs/괴산연습림_공기미생물_음이온_종합연계분석.xlsx")
wb.save(out_report)
print(f"Saved copy to {out_report} successfully!")
