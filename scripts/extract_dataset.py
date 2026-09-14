import sys
from pathlib import Path
import pandas as pd
import numpy as np
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

data_dir = Path(r"C:\Users\user\orca\workspaces\괴산 공기미생물\음이온-2\data\dated_summaries_20260911")

files = {
    "2026-04-15": data_dir / "음이온_260415_날짜포함_종합_시작시간수정.xlsx",
    "2026-04-29": data_dir / "음이온_260429_날짜포함_종합.xlsx",
    "2026-07-27": data_dir / "음이온_260727_날짜포함_종합.xlsx",
    "2026-08-13": data_dir / "음이온_260813_날짜포함_종합.xlsx",
}

site_names = [
    "리기다소나무림", "잣나무림 a", "버드나무림 a", "밭(대조구)", "주차장(대조구)",
    "소나무림", "일본잎갈나무림 a", "버드나무림 b", "일본잎갈나무림 b", "잣나무림 b"
]

site_types = {
    "리기다소나무림": "소나무류(리기다)",
    "잣나무림 a": "잣나무류",
    "잣나무림 b": "잣나무류",
    "버드나무림 a": "버드나무류",
    "버드나무림 b": "버드나무류",
    "소나무림": "소나무류",
    "일본잎갈나무림 a": "낙엽송(일본잎갈나무)",
    "일본잎갈나무림 b": "낙엽송(일본잎갈나무)",
    "밭(대조구)": "대조구(밭)",
    "주차장(대조구)": "대조구(주차장)"
}

is_forest = {
    "리기다소나무림": True,
    "잣나무림 a": True,
    "잣나무림 b": True,
    "버드나무림 a": True,
    "버드나무림 b": True,
    "소나무림": True,
    "일본잎갈나무림 a": True,
    "일본잎갈나무림 b": True,
    "밭(대조구)": False,
    "주차장(대조구)": False
}

records = []

for date_str, file_path in files.items():
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb["종합"]
    
    for row in range(2, 902):
        for s_idx in range(1, 11):
            val = ws.cell(row, s_idx).value
            ts = ws.cell(row, s_idx + 10).value
            
            is_num = isinstance(val, (int, float))
            num_val = float(val) if is_num else np.nan
            
            sname = site_names[s_idx - 1]
            records.append({
                "date": date_str,
                "site_id": s_idx,
                "site_name": sname,
                "site_type": site_types[sname],
                "is_forest": is_forest[sname],
                "second_offset": row - 2,
                "minute_bucket": (row - 2) // 60 + 1,
                "timestamp": str(ts),
                "ion_raw": str(val) if val is not None else "",
                "ion_value": num_val,
                "is_valid": is_num
            })
    wb.close()

df = pd.DataFrame(records)
out_csv = Path(r"C:\Users\user\orca\workspaces\괴산 공기미생물\음이온-2\data\ion_dataset_all.csv")
df.to_csv(out_csv, index=False, encoding='utf-8-sig')
print(f"Successfully extracted {len(df)} records to {out_csv}")
