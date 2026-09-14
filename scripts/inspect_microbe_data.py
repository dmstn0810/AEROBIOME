import sys
from pathlib import Path
import openpyxl
import pandas as pd

sys.stdout.reconfigure(encoding='utf-8')

file_path = Path("괴산연습림 전체데이터.xlsx")
wb = openpyxl.load_workbook(file_path, data_only=True)
print("Sheet names:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n--- Sheet: {sname} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
    for r in range(1, min(ws.max_row + 1, 10)):
        row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 25))]
        print(f"R{r}: {row_vals}")

wb.close()
