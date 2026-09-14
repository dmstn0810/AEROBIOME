import sys
from pathlib import Path
import pandas as pd
import numpy as np
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

data_dir = Path(r"C:\Users\user\orca\workspaces\괴산 공기미생물\음이온-2\data\dated_summaries_20260911")

files = {
    "2026-04-15": data_dir / "음이온_260415_날짜포함_종합_시작시간수정.xlsx",
    "2026-04-29": data_dir / "음이온_260429_날짜포함_종합.xlsx",
    "2026-07-27": data_dir / "음이온_260727_날짜포함_종합.xlsx",
    "2026-08-13": data_dir / "음이온_260813_날짜포함_종합.xlsx",
}

site_names = [
    "리기다소나무림", "잣나무림 a", "버드나무림 a", "밭(대조구)", "주차장(대조구)",
    "소나무림", "일본잎갈나무림 a", "버드나무림 b", "일본잎갈나무림 b", "잣나무림 b"
]

records = []

for date_str, file_path in files.items():
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb["종합"]
    
    # Check headers
    # Col 1..10: ion count
    # Col 11..20: timestamp
    # Col 21: survey date
    for row in range(2, 902): # 900 rows (row 2 to 901)
        for s_idx in range(1, 11):
            val = ws.cell(row, s_idx).value
            ts = ws.cell(row, s_idx + 10).value
            survey_date = ws.cell(row, 21).value
            
            # numeric check
            is_num = isinstance(val, (int, float))
            num_val = float(val) if is_num else np.nan
            
            records.append({
                "date": date_str,
                "site_id": s_idx,
                "site_name": site_names[s_idx - 1],
                "second_offset": row - 2, # 0 to 899
                "timestamp": ts,
                "ion_raw": val,
                "ion_value": num_val,
                "is_valid": is_num
            })
    wb.close()

df = pd.DataFrame(records)
print(f"Total records loaded: {len(df)}")
print(f"Valid numeric records: {df['is_valid'].sum()} / {len(df)}")
print(f"Invalid/non-numeric: {(~df['is_valid']).sum()}")

# Print non-numeric entries if any
invalid_rows = df[~df['is_valid']]
if not invalid_rows.empty:
    print("\nNon-numeric entries:")
    print(invalid_rows[["date", "site_id", "site_name", "second_offset", "timestamp", "ion_raw"]])

# Overall summary by date and site
summary = df.groupby(["date", "site_id", "site_name"]).agg(
    count=("ion_value", "count"),
    mean=("ion_value", "mean"),
    std=("ion_value", "std"),
    median=("ion_value", "median"),
    min=("ion_value", "min"),
    max=("ion_value", "max"),
    start_time=("timestamp", "min"),
    end_time=("timestamp", "max")
).reset_index()

print("\n--- Summary Table ---")
print(summary.to_string())
