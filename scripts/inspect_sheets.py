import sys
import openpyxl
from pathlib import Path
import pandas as pd

sys.stdout.reconfigure(encoding='utf-8')

data_dir = Path(r"C:\Users\user\orca\workspaces\괴산 공기미생물\음이온-2\data\dated_summaries_20260911")

files = [
    data_dir / "음이온_260415_날짜포함_종합_시작시간수정.xlsx",
    data_dir / "음이온_260429_날짜포함_종합.xlsx",
    data_dir / "음이온_260727_날짜포함_종합.xlsx",
    data_dir / "음이온_260813_날짜포함_종합.xlsx",
]

for file in files:
    print(f"=== File: {file.name} ===")
    wb = openpyxl.load_workbook(file, data_only=True)
    print("Sheets:", wb.sheetnames)
    
    # Check '자료구성'
    if '자료구성' in wb.sheetnames:
        ws_info = wb['자료구성']
        print("--- [자료구성] ---")
        for r in range(1, min(ws_info.max_row + 1, 25)):
            row_vals = [ws_info.cell(r, c).value for c in range(1, 10)]
            if any(row_vals):
                print(f"R{r}: {row_vals}")
                
    # Check '종합' sheet header
    if '종합' in wb.sheetnames:
        ws_sum = wb['종합']
        headers = [ws_sum.cell(1, c).value for c in range(1, 25)]
        print("--- [종합] Headers ---")
        print(headers)
        row2 = [ws_sum.cell(2, c).value for c in range(1, 25)]
        print("--- [종합] Row 2 Sample ---")
        print(row2)
        print(f"Total rows in 종합: {ws_sum.max_row}")
    
    wb.close()
    print("\n" + "="*50 + "\n")
